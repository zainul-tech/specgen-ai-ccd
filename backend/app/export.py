"""
Document export — REQ-EXPORT-01 (PDF), REQ-EXPORT-02 (DOCX), REQ-EXPORT-03
(preserve section structure and requirement identifiers).

Both builders take the same ORM Project object and walk it in the same section
order, so a PDF and a DOCX of the same project always have matching structure.
"""
import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, ListFlowable, ListItem,
)
from reportlab.lib import colors

NAVY = "1F3864"


# ---------------------------------------------------------------- shared content model

def _sections(project) -> dict:
    """Pulls everything the exported document needs out of the ORM project once,
    so both builders read from the same plain-data structure instead of each
    poking at SQLAlchemy relationships independently (and risking drifting apart)."""
    return {
        "idea": project.idea_text,
        "classification": {
            "type": project.classification_type, "domain": project.classification_domain,
            "users": project.classification_users,
        },
        "status": project.status,
        "quality_score": project.quality_score,
        "requirements": [
            {"req_id": r.req_id, "text": r.text, "priority": r.priority} for r in project.requirements
        ],
        "stories": [{"role": s.role, "goal": s.goal, "benefit": s.benefit} for s in project.user_stories],
        "use_cases": [
            {"name": u.name, "actor": u.actor, "precondition": u.precondition,
             "main_flow": u.main_flow, "postcondition": u.postcondition}
            for u in project.use_cases
        ],
        "risks": [
            {"description": r.description, "likelihood": r.likelihood, "impact": r.impact,
             "mitigation": r.mitigation}
            for r in project.risks
        ],
        "cost_estimates": [
            {"area": c.area, "effort_level": c.effort_level, "notes": c.notes} for c in project.cost_estimates
        ],
        "sprints": [
            {"sprint_no": s.sprint_no, "duration_weeks": s.duration_weeks,
             "stories": [f"As a {st.role}, I want {st.goal}." for st in s.stories]}
            for s in project.sprints
        ],
        "test_cases": [
            {"feature": t.feature, "input": t.input_condition, "expected": t.expected_result}
            for t in project.test_cases
        ],
    }


# ---------------------------------------------------------------- DOCX

def build_docx(project) -> bytes:
    d = _sections(project)
    doc = Document()

    title = doc.add_heading("SpecGen AI — Project Export", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Idea", level=1)
    doc.add_paragraph(d["idea"])

    doc.add_heading("2. Classification", level=1)
    cls = d["classification"]
    doc.add_paragraph(f"Type: {cls['type'] or '—'}")
    doc.add_paragraph(f"Domain: {cls['domain'] or '—'}")
    doc.add_paragraph(f"Users: {cls['users'] or '—'}")
    doc.add_paragraph(f"Status: {d['status']}" + (f"  ·  Quality score: {d['quality_score']}/100" if d["quality_score"] is not None else ""))

    doc.add_heading("3. Functional Requirements", level=1)
    if not d["requirements"]:
        doc.add_paragraph("None.")
    for r in d["requirements"]:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(f"{r['req_id']}: ")
        run.bold = True
        p.add_run(f"{r['text']} [{r['priority']}]")

    doc.add_heading("4. User Stories", level=1)
    if not d["stories"]:
        doc.add_paragraph("None.")
    for s in d["stories"]:
        doc.add_paragraph(f"As a {s['role']}, I want {s['goal']}, so that {s['benefit']}.", style="List Bullet")

    doc.add_heading("5. Use Cases", level=1)
    if not d["use_cases"]:
        doc.add_paragraph("None.")
    for u in d["use_cases"]:
        doc.add_heading(f"{u['name']} ({u['actor']})", level=2)
        if u["precondition"]:
            doc.add_paragraph(f"Precondition: {u['precondition']}")
        doc.add_paragraph(f"Main flow: {u['main_flow']}")
        if u["postcondition"]:
            doc.add_paragraph(f"Postcondition: {u['postcondition']}")

    doc.add_heading("6. Risks", level=1)
    if d["risks"]:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Risk", "Likelihood", "Impact", "Mitigation"]):
            hdr[i].text = h
        for r in d["risks"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = r["description"], r["likelihood"], r["impact"]
            row[3].text = r["mitigation"] or "—"
    else:
        doc.add_paragraph("None.")

    doc.add_heading("7. Cost Estimate", level=1)
    if d["cost_estimates"]:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Area", "Effort", "Notes"]):
            hdr[i].text = h
        for c in d["cost_estimates"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = c["area"], c["effort_level"], c["notes"] or "—"
    else:
        doc.add_paragraph("None.")

    doc.add_heading("8. Sprint Plan", level=1)
    if d["sprints"]:
        for s in d["sprints"]:
            doc.add_heading(f"Sprint {s['sprint_no']} ({s['duration_weeks']} weeks)", level=2)
            if s["stories"]:
                for story in s["stories"]:
                    doc.add_paragraph(story, style="List Bullet")
            else:
                doc.add_paragraph("No stories assigned.")
    else:
        doc.add_paragraph("No sprint plan generated.")

    doc.add_heading("9. Test Cases", level=1)
    if d["test_cases"]:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Feature", "Input", "Expected Result"]):
            hdr[i].text = h
        for t in d["test_cases"]:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = t["feature"], t["input"], t["expected"]
    else:
        doc.add_paragraph("None.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def build_pdf(project) -> bytes:
    d = _sections(project)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=0.75 * inch, bottomMargin=0.75 * inch,
                             leftMargin=0.75 * inch, rightMargin=0.75 * inch)

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=14)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=8)
    body = ParagraphStyle("Body", parent=styles["Normal"], alignment=TA_LEFT, spaceAfter=4, leading=14)
    title_style = ParagraphStyle("Title", parent=styles["Title"], textColor=colors.HexColor(f"#{NAVY}"))

    story = [Paragraph("SpecGen AI — Project Export", title_style), Spacer(1, 16)]

    story.append(Paragraph("1. Idea", h1))
    story.append(Paragraph(d["idea"], body))

    story.append(Paragraph("2. Classification", h1))
    cls = d["classification"]
    story.append(Paragraph(f"Type: {cls['type'] or '—'}", body))
    story.append(Paragraph(f"Domain: {cls['domain'] or '—'}", body))
    story.append(Paragraph(f"Users: {cls['users'] or '—'}", body))
    status_line = f"Status: {d['status']}"
    if d["quality_score"] is not None:
        status_line += f"  ·  Quality score: {d['quality_score']}/100"
    story.append(Paragraph(status_line, body))

    story.append(Paragraph("3. Functional Requirements", h1))
    if d["requirements"]:
        items = [ListItem(Paragraph(f"<b>{r['req_id']}:</b> {r['text']} [{r['priority']}]", body))
                 for r in d["requirements"]]
        story.append(ListFlowable(items, bulletType="bullet"))
    else:
        story.append(Paragraph("None.", body))

    story.append(Paragraph("4. User Stories", h1))
    if d["stories"]:
        items = [ListItem(Paragraph(f"As a {s['role']}, I want {s['goal']}, so that {s['benefit']}.", body))
                 for s in d["stories"]]
        story.append(ListFlowable(items, bulletType="bullet"))
    else:
        story.append(Paragraph("None.", body))

    story.append(Paragraph("5. Use Cases", h1))
    if d["use_cases"]:
        for u in d["use_cases"]:
            story.append(Paragraph(f"{u['name']} ({u['actor']})", h2))
            if u["precondition"]:
                story.append(Paragraph(f"Precondition: {u['precondition']}", body))
            story.append(Paragraph(f"Main flow: {u['main_flow']}", body))
            if u["postcondition"]:
                story.append(Paragraph(f"Postcondition: {u['postcondition']}", body))
    else:
        story.append(Paragraph("None.", body))

    story.append(Paragraph("6. Risks", h1))
    if d["risks"]:
        rows = [["Risk", "Likelihood", "Impact", "Mitigation"]]
        rows += [[r["description"], r["likelihood"], r["impact"], r["mitigation"] or "—"] for r in d["risks"]]
        story.append(_table(rows))
    else:
        story.append(Paragraph("None.", body))

    story.append(Paragraph("7. Cost Estimate", h1))
    if d["cost_estimates"]:
        rows = [["Area", "Effort", "Notes"]]
        rows += [[c["area"], c["effort_level"], c["notes"] or "—"] for c in d["cost_estimates"]]
        story.append(_table(rows))
    else:
        story.append(Paragraph("None.", body))

    story.append(Paragraph("8. Sprint Plan", h1))
    if d["sprints"]:
        for s in d["sprints"]:
            story.append(Paragraph(f"Sprint {s['sprint_no']} ({s['duration_weeks']} weeks)", h2))
            if s["stories"]:
                story.append(ListFlowable([ListItem(Paragraph(st, body)) for st in s["stories"]], bulletType="bullet"))
            else:
                story.append(Paragraph("No stories assigned.", body))
    else:
        story.append(Paragraph("No sprint plan generated.", body))

    story.append(Paragraph("9. Test Cases", h1))
    if d["test_cases"]:
        rows = [["Feature", "Input", "Expected Result"]]
        rows += [[t["feature"], t["input"], t["expected"]] for t in d["test_cases"]]
        story.append(_table(rows))
    else:
        story.append(Paragraph("None.", body))

    doc.build(story)
    return buf.getvalue()


def _table(rows) -> Table:
    t = Table(rows, hAlign="LEFT", repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t
