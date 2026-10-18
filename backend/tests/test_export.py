"""
Export tests (REQ-EXPORT-01/02/03). Deliberately read the generated files back
with real parsers (pypdf, python-docx) rather than just checking status codes —
a 200 response with garbage bytes would pass a status-only test and still be
useless to a user.
"""
import io
import pypdf
from docx import Document


def _get(client, pid, fmt, headers):
    return client.get(f"/projects/{pid}/export?format={fmt}", headers=headers)


# ------------------------------------------------------------------ access control (shared with every other route)

def test_export_requires_auth(client, project):
    assert _get(client, project["id"], "pdf", {}).status_code == 401


def test_export_requires_ownership(client, project, other_headers):
    assert _get(client, project["id"], "pdf", other_headers).status_code == 404
    assert _get(client, project["id"], "docx", other_headers).status_code == 404


def test_invalid_format_rejected(client, project, auth_headers):
    r = _get(client, project["id"], "xlsx", auth_headers)
    assert r.status_code == 400


# ------------------------------------------------------------------ PDF: real parser, real content

def test_pdf_is_a_real_pdf_with_correct_headers(client, project, auth_headers):
    r = _get(client, project["id"], "pdf", auth_headers)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert "specgen-ai-project" in r.headers["content-disposition"]
    assert r.content[:4] == b"%PDF"


def test_pdf_text_contains_requirement_ids_and_idea(client, project, auth_headers):
    r = _get(client, project["id"], "pdf", auth_headers)
    reader = pypdf.PdfReader(io.BytesIO(r.content))
    text = "\n".join(page.extract_text() for page in reader.pages)

    for req in project["requirements"]:
        assert req["req_id"] in text, f"{req['req_id']} missing from exported PDF text"
    assert "food ordering" in text  # from the idea text in the canned draft
    assert "Payment gateway failure" in text  # risk description


def test_pdf_reflects_edits_made_before_export(client, project, auth_headers):
    """REQ-EXPORT-03 is only meaningful if export reads current state, not a cached
    copy of what generation originally produced."""
    pid = project["id"]
    req_pk = project["requirements"][0]["id"]
    client.patch(f"/projects/{pid}/requirements/{req_pk}",
                 json={"text": "The system shall load pages within 2 seconds, edited for export test."},
                 headers=auth_headers)

    r = _get(client, pid, "pdf", auth_headers)
    reader = pypdf.PdfReader(io.BytesIO(r.content))
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "edited for export test" in text


def test_pdf_handles_empty_optional_sections_without_crashing(client, auth_headers, monkeypatch):
    """A project where generation returned no risks/use-cases/etc. should still export
    cleanly, not throw when a section has nothing to render."""
    from app import ai_engine
    monkeypatch.setattr(ai_engine, "generate_draft", lambda idea: {
        "classification": {}, "functionalRequirements": [], "userStories": [],
        "useCases": [], "risks": [], "testCases": [], "costEstimate": [],
    })
    pid = client.post("/projects", json={"idea_text": "A minimal test idea with no generated content at all."},
                      headers=auth_headers).json()["id"]
    r = _get(client, pid, "pdf", auth_headers)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"


# ------------------------------------------------------------------ DOCX: real parser, real content

def test_docx_is_a_real_docx_with_correct_headers(client, project, auth_headers):
    r = _get(client, project["id"], "docx", auth_headers)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert r.content[:2] == b"PK"  # docx is a zip archive


def test_docx_text_contains_requirement_ids_and_structure(client, project, auth_headers):
    r = _get(client, project["id"], "docx", auth_headers)
    doc = Document(io.BytesIO(r.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    for req in project["requirements"]:
        assert req["req_id"] in full_text

    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    # REQ-EXPORT-03: section structure preserved, in order.
    expected_order = ["Idea", "Classification", "Functional Requirements", "User Stories",
                       "Use Cases", "Risks", "Cost Estimate", "Sprint Plan", "Test Cases"]
    found_order = [h for h in headings if any(e in h for e in expected_order)]
    assert len(found_order) == 9
    for i, label in enumerate(expected_order):
        assert label in found_order[i]


def test_docx_risk_table_has_mitigation_column(client, project, auth_headers):
    r = _get(client, project["id"], "docx", auth_headers)
    doc = Document(io.BytesIO(r.content))
    assert len(doc.tables) >= 1
    risk_table = doc.tables[0]
    header_cells = [c.text for c in risk_table.rows[0].cells]
    assert header_cells == ["Risk", "Likelihood", "Impact", "Mitigation"]
    assert risk_table.rows[1].cells[3].text  # mitigation text present, not blank


def test_docx_includes_sprint_plan_when_generated(client, project, auth_headers):
    pid = project["id"]
    client.post(f"/projects/{pid}/sprint-plan", json={"sprint_count": 2, "duration_weeks": 1}, headers=auth_headers)
    r = _get(client, pid, "docx", auth_headers)
    doc = Document(io.BytesIO(r.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Sprint 1" in full_text and "Sprint 2" in full_text


def test_docx_omits_sprint_plan_gracefully_when_not_generated(client, project, auth_headers):
    r = _get(client, project["id"], "docx", auth_headers)
    doc = Document(io.BytesIO(r.content))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "No sprint plan generated." in full_text
